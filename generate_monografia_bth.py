#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador Automatizado de la Monografía Técnica BTH 2026
Especialidad: Sistemas Informáticos - Instituto Americano "AMERINST"
Proyecto: "SISTEMA WEB DE COMERCIO ELECTRÓNICO Y DESPACHO LOGÍSTICO EN TIEMPO REAL
CON CONTROL DE INVENTARIOS APLICANDO CÓDIGOS QR Y GEORREFERENCIACIÓN PARA LA EMPRESA 'BURGER 24/7'"

Genera:
1. MONOGRAFIA_BTH_2026_BURGER247.docx (Documento Word con formato formal BTH, Arial 12pt,
   márgenes de 1 pulgada, carátula, preliminares, tablas estilizadas y Marco Teórico >= 15 páginas).
2. Actualización directa de C:\\Users\\natal\\Downloads\\Formato Documento BTH-2026.docx
3. MONOGRAFIA_BTH_2026_BURGER247.html (Visor web interactivo con renderizado de diagramas Mermaid.js).
"""

import os
import shutil
import re
import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import bth_monografia_content_part1 as part1
import bth_monografia_content_marco_teorico as part2
import bth_monografia_content_part3 as part3

ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(ROOT_DIR, "MONOGRAFIA_BTH_2026_BURGER247.docx")
OUT_HTML = os.path.join(ROOT_DIR, "MONOGRAFIA_BTH_2026_BURGER247.html")
DOWNLOADS_DOCX = r"C:\Users\natal\Downloads\Formato Documento BTH-2026.docx"

# ==============================================================================
# HELPERS DE MAQUETACIÓN OPENXML
# ==============================================================================

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Establece padding interno en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Establece bordes sutiles y elegantes a una tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_footer_page_number(run):
    """Inserta el campo dinámico PAGE en el run del pie de página."""
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_inline_formatted_text(paragraph, text):
    """
    Parsea markdown simple (**negrita**, *cursiva*, `código`, fórmulas $...$)
    y lo inserta como runs con el formato adecuado.
    """
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|“.*?”|«.*?»)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(180, 83, 9)
        elif part.startswith('$') and part.endswith('$') and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Cambria Math'
            r.italic = True
            r.font.color.rgb = RGBColor(30, 58, 138)
        else:
            paragraph.add_run(part)

def render_table_in_doc(doc, table_data):
    """Renderiza una tabla de datos estilizada con encabezado azul marino y filas alternadas."""
    headers = table_data["columnas"]
    rows = table_data["filas"]
    
    # Párrafo de título de la tabla
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(table_data["titulo"])
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)

    # Encabezados
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")  # Navy Blue
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=180, right=180)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Filas de datos
    for r_idx, row_values in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=160, right=160)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def render_code_block(doc, code_text, label="CÓDIGO FUENTE"):
    """Inserta un bloque de código estilizado con fondo sombreado y fuente Consolas."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Borde izquierdo grueso azul/ámbar
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_lbl = p.add_run(f"[{label}]\n")
    r_lbl.font.name = "Consolas"
    r_lbl.font.size = Pt(8.5)
    r_lbl.bold = True
    r_lbl.font.color.rgb = RGBColor(37, 99, 235)
    
    r_code = p.add_run(code_text.strip())
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def render_mermaid_diagram(doc, diag_info):
    """Inserta un diagrama Mermaid en bloque estilizado con descripción técnica."""
    # Título del diagrama
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(f"Figura {diag_info['numero']}: {diag_info['titulo']}")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    # Descripción
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_desc.add_run(diag_info["descripcion"])
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10)
    r_desc.italic = True
    r_desc.font.color.rgb = RGBColor(71, 85, 105)
    
    # Bloque de código sintaxis Mermaid
    render_code_block(doc, diag_info["mermaid"], label=f"DIAGRAMA MERMAID {diag_info['numero']} - SINTAXIS OFICIAL")

# ==============================================================================
# CONSTRUCTOR PRINCIPAL DE LA MONOGRAFÍA
# ==============================================================================

def build_monografia_docx():
    print("[1/5] Inicializando documento Word con maquetación oficial BTH 2026...")
    doc = Document()
    
    # 1. Configuración de página: Formato Carta, márgenes de 1.0 pulgada (72 pt / 1440 dxa)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Encabezado (Páginas subsiguientes)
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("MONOGRAFÍA BTH 2026 | SISTEMAS INFORMÁTICOS - “AMERINST”")
        r_hdr.font.name = "Arial"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)
        
        # Pie de página (Páginas subsiguientes)
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Página ")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(9.5)
        r_ftr.font.color.rgb = RGBColor(100, 116, 139)
        add_footer_page_number(r_ftr)
        r_ftr2 = p_ftr.add_run(" | Burger 24/7 E-Commerce")
        r_ftr2.font.name = "Arial"
        r_ftr2.font.size = Pt(9.5)
        r_ftr2.font.color.rgb = RGBColor(148, 163, 184)

    # 2. Configurar estilo Normal a Arial 12pt Justificado
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(17, 24, 39)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(5)

    # ==========================================================================
    # CARÁTULA FORMAL BTH
    # ==========================================================================
    print("[2/5] Construyendo Carátula y Secciones Preliminares...")
    meta = part1.METADATA
    
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(24)
    p_inst.paragraph_format.space_after = Pt(2)
    r1 = p_inst.add_run(f"{meta['institucion']}\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(16)
    r1.bold = True
    r1.font.color.rgb = RGBColor(30, 41, 59)
    
    r2 = p_inst.add_run(f"{meta['colegio']}\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(18)
    r2.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)
    
    r3 = p_inst.add_run(f"{meta['especialidad']}\n")
    r3.font.name = "Arial"
    r3.font.size = Pt(13)
    r3.bold = True
    r3.font.color.rgb = RGBColor(71, 85, 105)
    
    r_sub = p_inst.add_run(f"{meta['subtitulo_bth']}\n\n\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(180, 83, 9)

    # Título de la Monografía
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(24)
    p_tit.paragraph_format.space_after = Pt(36)
    p_tit.paragraph_format.line_spacing = 1.25
    r_tit = p_tit.add_run(meta["titulo"])
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(16)
    r_tit.bold = True
    r_tit.font.color.rgb = RGBColor(15, 23, 42)

    # Datos de los postulantes y tutor
    p_post = doc.add_paragraph()
    p_post.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_post.paragraph_format.space_before = Pt(36)
    p_post.paragraph_format.space_after = Pt(4)
    post_txt = meta.get("postulantes", meta.get("postulante", "Nataly Gemio"))
    if not post_txt.startswith("Postulante"):
        post_label = "Postulantes" if (" y " in post_txt or "Equipo" in post_txt) else "Postulante"
        r_post = p_post.add_run(f"{post_label}: {post_txt}\n")
    else:
        r_post = p_post.add_run(f"{post_txt}\n")
    r_post.font.name = "Arial"
    r_post.font.size = Pt(12)
    r_post.bold = True
    
    r_tut = p_post.add_run(f"Tutor: {meta['tutor']}\n")
    r_tut.font.name = "Arial"
    r_tut.font.size = Pt(12)
    r_tut.bold = True
    
    r_cur = p_post.add_run(f"Curso: {meta['curso']}\n\n\n\n")
    r_cur.font.name = "Arial"
    r_cur.font.size = Pt(12)
    r_cur.bold = True

    # Pie de carátula: Lugar y Fecha
    p_pie = doc.add_paragraph()
    p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pie.paragraph_format.space_before = Pt(48)
    r_lugar = p_pie.add_run(f"{meta['lugar']}\n{meta['gestion']}")
    r_lugar.font.name = "Arial"
    r_lugar.font.size = Pt(12)
    r_lugar.bold = True
    r_lugar.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # ==========================================================================
    # SECCIONES PRELIMINARES
    # ==========================================================================
    preliminares = [
        ("DEDICATORIA", part1.PRELIMINARES["dedicatoria"]),
        ("AGRADECIMIENTO", part1.PRELIMINARES["agradecimiento"]),
        ("CITA BÍBLICA", part1.PRELIMINARES["cita_biblica"]),
        ("ABSTRACT", part1.PRELIMINARES["abstract"]),
        ("RESUMEN", part1.PRELIMINARES["resumen"])
    ]

    for titulo_sec, cuerpo_sec in preliminares:
        p_h = doc.add_paragraph()
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h.paragraph_format.space_before = Pt(18)
        p_h.paragraph_format.space_after = Pt(18)
        p_h.paragraph_format.keep_with_next = True
        r_h = p_h.add_run(titulo_sec)
        r_h.font.name = "Arial"
        r_h.font.size = Pt(14)
        r_h.bold = True
        r_h.font.color.rgb = RGBColor(30, 58, 138)

        for parrafo in cuerpo_sec.split("\n\n"):
            if not parrafo.strip():
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(8)
            add_inline_formatted_text(p, parrafo.strip())

        doc.add_page_break()

    # ==========================================================================
    # ÍNDICE GENERAL ESTRUCTURADO
    # ==========================================================================
    p_idx_tit = doc.add_paragraph()
    p_idx_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_idx_tit.paragraph_format.space_before = Pt(18)
    p_idx_tit.paragraph_format.space_after = Pt(18)
    p_idx_tit.paragraph_format.keep_with_next = True
    r_idx_tit = p_idx_tit.add_run("ÍNDICE GENERAL")
    r_idx_tit.font.name = "Arial"
    r_idx_tit.font.size = Pt(14)
    r_idx_tit.bold = True
    r_idx_tit.font.color.rgb = RGBColor(30, 58, 138)

    indice_items = [
        ("DEDICATORIA", "ii"),
        ("AGRADECIMIENTO", "iii"),
        ("CITA BÍBLICA", "iv"),
        ("ABSTRACT", "v"),
        ("RESUMEN", "vi"),
        ("ÍNDICE GENERAL", "vii"),
        ("I. INTRODUCCIÓN", "1"),
        ("II. PLANTEAMIENTO DEL PROBLEMA", "3"),
        ("    2.1 PROBLEMA PRINCIPAL", "3"),
        ("    2.2 PROBLEMAS SECUNDARIOS", "4"),
        ("III. JUSTIFICACIÓN", "5"),
        ("    3.1 Justificación Técnica", "5"),
        ("    3.2 Justificación Económica", "6"),
        ("    3.3 Justificación Social", "7"),
        ("IV. ALCANCES Y DELIMITACIONES", "8"),
        ("    4.1 DESTINATARIOS", "8"),
        ("    4.2 Delimitación Espacio-Temporal y Tecnológica", "9"),
        ("V. LOCALIZACIÓN O UBICACIÓN", "10"),
        ("VI. OBJETIVOS", "11"),
        ("    6.1 OBJETIVO GENERAL", "11"),
        ("    6.2 OBJETIVOS ESPECÍFICOS", "11"),
        ("VII. MARCO TEÓRICO (CONCEPTOS Y SUSTENTO TÉCNICO)", "12"),
        ("    7.1 SUSTENTO LEGAL", "12"),
        ("        7.1.1 CPE Bolivia (Arts. 103, 47, 75)", "12"),
        ("        7.1.2 Ley N° 164 de Telecomunicaciones y Comercio Electrónico", "13"),
        ("        7.1.3 Decreto Supremo N° 1793 (Software Libre y Estándares Abiertos)", "14"),
        ("        7.1.4 Ley N° 453 de Derechos de las Usuarias y Consumidores", "15"),
        ("        7.1.5 Código de Comercio de Bolivia (Actos Mercantiles Digitales)", "16"),
        ("        7.1.6 Normativa ASFI para Pagos Móviles y Códigos Simple QR", "17"),
        ("        7.1.7 Ley de Educación N° 070 y Lineamientos BTH", "18"),
        ("        7.1.8 Estándares Internacionales ISO/IEC 25010 e ISO/IEC 27001", "19"),
        ("    7.2 DESARROLLO DEL MARCO TEÓRICO (CONCEPTOS DEL SISTEMA)", "20"),
        ("        7.2.1 Arquitectura Cliente-Servidor y Conexiones Web (HTTP, REST, JSON)", "20"),
        ("        7.2.2 Tecnologías Frontend: HTML5, CSS3 y Vanilla JavaScript ES6+", "22"),
        ("        7.2.3 Tecnologías Backend: PHP 8.2 y Abstracción de Datos con PDO", "24"),
        ("        7.2.4 Tecnologías de Computación Matemática: Python 3.11 en Logística", "26"),
        ("        7.2.5 Sistemas Gestores de BD: MySQL/MariaDB, InnoDB y 3FN", "28"),
        ("        7.2.6 Control de Concurrencia: Bloqueo Pesimista (SELECT ... FOR UPDATE)", "31"),
        ("        7.2.7 Criptografía y Seguridad Web: Bcrypt, Tokens JWT y OWASP", "33"),
        ("        7.2.8 Sistemas de Información Geográfica: Leaflet, OSM y Fórmula Haversine", "35"),
        ("        7.2.9 Pasarelas de Pago Digitales y Códigos QR: EMVCo y Simple QR", "38"),
        ("        7.2.10 Control de Versiones y Trabajo en Equipo: Git y GitHub", "40"),
        ("        7.2.11 Auditoría de Sistemas, Logs Inmutables y Sobres BMAD", "42"),
        ("        7.2.12 Metodologías de Calidad de Software y Verificación E2E", "44"),
        ("VIII. MARCO PROCEDIMENTAL", "46"),
        ("    8.1 PROPUESTA DE INNOVACIÓN", "46"),
        ("    8.2 RESULTADOS ESPERADOS", "47"),
        ("IX. METODOLOGÍA", "48"),
        ("    9.1 TIPO DE INVESTIGACIÓN", "48"),
        ("    9.2 TÉCNICAS E INSTRUMENTOS DE RECOLECCIÓN DE DATOS", "49"),
        ("X. METODOLOGÍA DE DESARROLLO DE SOFTWARE", "50"),
        ("    Explicación de la Metodología en Cascada Clásica (Waterfall)", "50"),
        ("    10.1 ANÁLISIS (Requerimientos RF y RNF, Casos de Uso)", "51"),
        ("    10.2 DISEÑO (Modelado C4, DER Físico 3FN, Secuencia, Estados)", "53"),
        ("    10.3 IMPLEMENTACIÓN (Codificación Frontend/Backend)", "60"),
        ("    10.4 VERIFICACIÓN (Pruebas Unitarias, Integración y E2E)", "62"),
        ("    10.5 MANTENIMIENTO (Políticas de Backup y Seguridad)", "64"),
        ("    10.6 CRONOGRAMA DE ACTIVIDADES (Diagrama de Gantt Cascada 2026)", "65"),
        ("    10.7 RECURSOS (Materiales, 2 Estudiantes Devs y Presupuesto)", "66"),
        ("XI. ARTICULACIÓN CON CAMPOS Y ÁREAS DE SABERES Y CONOCIMIENTOS", "68"),
        ("XII. CONCLUSIONES Y RECOMENDACIONES", "71"),
        ("XIII. PROYECTO DE VIDA", "73"),
        ("XIV. BIBLIOGRAFÍA (Normas APA 7ma Edición)", "75"),
        ("ANEXOS", "78"),
        ("    ANEXO A: Diccionario de Datos Físico de MySQL (Motor InnoDB)", "78"),
        ("    ANEXO B: Catálogo de Endpoints de la API REST", "81"),
        ("    ANEXO C: Algoritmos Fundamentales en Código Fuente", "83")
    ]

    p_idx_table = doc.add_table(rows=len(indice_items), cols=2)
    p_idx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(p_idx_table, color="FFFFFF") # Sin bordes visibles

    for i, (item_text, pag_num) in enumerate(indice_items):
        cell_item, cell_pag = p_idx_table.rows[i].cells
        cell_item.text = item_text
        set_cell_margins(cell_item, top=30, bottom=30, left=50, right=50)
        p_i = cell_item.paragraphs[0]
        p_i.paragraph_format.line_spacing = 1.0
        p_i.paragraph_format.space_after = Pt(2)
        r_i = p_i.runs[0]
        r_i.font.name = "Arial"
        r_i.font.size = Pt(10)
        if not item_text.startswith("    "):
            r_i.bold = True
            r_i.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r_i.font.color.rgb = RGBColor(51, 65, 85)

        cell_pag.text = pag_num
        set_cell_margins(cell_pag, top=30, bottom=30, left=50, right=50)
        p_p = cell_pag.paragraphs[0]
        p_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_p.paragraph_format.line_spacing = 1.0
        p_p.paragraph_format.space_after = Pt(2)
        r_p = p_p.runs[0]
        r_p.font.name = "Arial"
        r_p.font.size = Pt(10)
        r_p.bold = not item_text.startswith("    ")
        r_p.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # ==========================================================================
    # CAPÍTULOS I AL VI
    # ==========================================================================
    print("[3/5] Redactando Capítulos I al VI...")
    
    # Cap I: Introducción
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(16)
    p_c1.paragraph_format.space_after = Pt(12)
    p_c1.paragraph_format.keep_with_next = True
    r_c1 = p_c1.add_run(part1.CAPITULO_I["titulo"])
    r_c1.font.name = "Arial"
    r_c1.font.size = Pt(14)
    r_c1.bold = True
    r_c1.font.color.rgb = RGBColor(30, 58, 138)
    
    for parrafo in part1.CAPITULO_I["contenido"].split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p, parrafo.strip())
    
    doc.add_page_break()

    # Cap II: Planteamiento del Problema
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(16)
    p_c2.paragraph_format.space_after = Pt(12)
    p_c2.paragraph_format.keep_with_next = True
    r_c2 = p_c2.add_run(part1.CAPITULO_II["titulo"])
    r_c2.font.name = "Arial"
    r_c2.font.size = Pt(14)
    r_c2.bold = True
    r_c2.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part1.CAPITULO_II["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap III: Justificación
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.space_before = Pt(16)
    p_c3.paragraph_format.space_after = Pt(12)
    p_c3.paragraph_format.keep_with_next = True
    r_c3 = p_c3.add_run(part1.CAPITULO_III["titulo"])
    r_c3.font.name = "Arial"
    r_c3.font.size = Pt(14)
    r_c3.bold = True
    r_c3.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part1.CAPITULO_III["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap IV: Alcances y Delimitaciones
    p_c4 = doc.add_paragraph()
    p_c4.paragraph_format.space_before = Pt(16)
    p_c4.paragraph_format.space_after = Pt(12)
    p_c4.paragraph_format.keep_with_next = True
    r_c4 = p_c4.add_run(part1.CAPITULO_IV["titulo"])
    r_c4.font.name = "Arial"
    r_c4.font.size = Pt(14)
    r_c4.bold = True
    r_c4.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part1.CAPITULO_IV["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap V: Localización o Ubicación
    p_c5 = doc.add_paragraph()
    p_c5.paragraph_format.space_before = Pt(16)
    p_c5.paragraph_format.space_after = Pt(12)
    p_c5.paragraph_format.keep_with_next = True
    r_c5 = p_c5.add_run(part1.CAPITULO_V["titulo"])
    r_c5.font.name = "Arial"
    r_c5.font.size = Pt(14)
    r_c5.bold = True
    r_c5.font.color.rgb = RGBColor(30, 58, 138)

    for parrafo in part1.CAPITULO_V["contenido"].split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap VI: Objetivos
    p_c6 = doc.add_paragraph()
    p_c6.paragraph_format.space_before = Pt(16)
    p_c6.paragraph_format.space_after = Pt(12)
    p_c6.paragraph_format.keep_with_next = True
    r_c6 = p_c6.add_run(part1.CAPITULO_VI["titulo"])
    r_c6.font.name = "Arial"
    r_c6.font.size = Pt(14)
    r_c6.bold = True
    r_c6.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part1.CAPITULO_VI["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # ==========================================================================
    # CAPÍTULO VII: MARCO TEÓRICO (>= 15 HOJAS OBLIGATORIO)
    # ==========================================================================
    print("[4/5] Ensamblando Capítulo VII: MARCO TEÓRICO exhaustivo (15+ hojas garantizadas)...")
    mt = part2.CAPITULO_VII
    
    p_c7 = doc.add_paragraph()
    p_c7.paragraph_format.space_before = Pt(16)
    p_c7.paragraph_format.space_after = Pt(12)
    p_c7.paragraph_format.keep_with_next = True
    r_c7 = p_c7.add_run(mt["titulo"])
    r_c7.font.name = "Arial"
    r_c7.font.size = Pt(15)
    r_c7.bold = True
    r_c7.font.color.rgb = RGBColor(30, 58, 138)

    # 7.1 SUSTENTO LEGAL
    sec_7_1 = mt["seccion_7_1"]
    p_s71 = doc.add_paragraph()
    p_s71.paragraph_format.space_before = Pt(14)
    p_s71.paragraph_format.space_after = Pt(6)
    p_s71.paragraph_format.keep_with_next = True
    r_s71 = p_s71.add_run(sec_7_1["subtitulo"])
    r_s71.font.name = "Arial"
    r_s71.font.size = Pt(13)
    r_s71.bold = True
    r_s71.font.color.rgb = RGBColor(15, 23, 42)

    p_d71 = doc.add_paragraph()
    p_d71.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_formatted_text(p_d71, sec_7_1["descripcion"])

    for item in sec_7_1["items"]:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(10)
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.keep_with_next = True
        r_item = p_item.add_run(f"{item['codigo']} {item['nombre']}")
        r_item.font.name = "Arial"
        r_item.font.size = Pt(12)
        r_item.bold = True
        r_item.font.color.rgb = RGBColor(30, 41, 59)

        for parrafo in item["analisis"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    if "tabla_legal" in sec_7_1:
        render_table_in_doc(doc, sec_7_1["tabla_legal"])

    # 7.2 DESARROLLO DEL MARCO TEÓRICO
    sec_7_2 = mt["seccion_7_2"]
    p_s72 = doc.add_paragraph()
    p_s72.paragraph_format.space_before = Pt(16)
    p_s72.paragraph_format.space_after = Pt(6)
    p_s72.paragraph_format.keep_with_next = True
    r_s72 = p_s72.add_run(sec_7_2["subtitulo"])
    r_s72.font.name = "Arial"
    r_s72.font.size = Pt(13)
    r_s72.bold = True
    r_s72.font.color.rgb = RGBColor(15, 23, 42)

    p_d72 = doc.add_paragraph()
    p_d72.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_formatted_text(p_d72, sec_7_2["descripcion"])

    for tema in sec_7_2["temas"]:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(12)
        p_t.paragraph_format.space_after = Pt(4)
        p_t.paragraph_format.keep_with_next = True
        r_t = p_t.add_run(f"{tema['numero']} {tema['titulo']}")
        r_t.font.name = "Arial"
        r_t.font.size = Pt(12)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(30, 41, 59)

        for parrafo in tema["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

        if "tabla" in tema:
            render_table_in_doc(doc, tema["tabla"])

    doc.add_page_break()

    # ==========================================================================
    # CAPÍTULOS VIII Y IX
    # ==========================================================================
    # Cap VIII: Marco Procedimental
    p_c8 = doc.add_paragraph()
    p_c8.paragraph_format.space_before = Pt(16)
    p_c8.paragraph_format.space_after = Pt(12)
    p_c8.paragraph_format.keep_with_next = True
    r_c8 = p_c8.add_run(part3.CAPITULO_VIII["titulo"])
    r_c8.font.name = "Arial"
    r_c8.font.size = Pt(14)
    r_c8.bold = True
    r_c8.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part3.CAPITULO_VIII["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap IX: Metodología
    p_c9 = doc.add_paragraph()
    p_c9.paragraph_format.space_before = Pt(16)
    p_c9.paragraph_format.space_after = Pt(12)
    p_c9.paragraph_format.keep_with_next = True
    r_c9 = p_c9.add_run(part3.CAPITULO_IX["titulo"])
    r_c9.font.name = "Arial"
    r_c9.font.size = Pt(14)
    r_c9.bold = True
    r_c9.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part3.CAPITULO_IX["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # ==========================================================================
    # CAPÍTULO X: METODOLOGÍA DE DESARROLLO DE SOFTWARE (CON MERMAID)
    # ==========================================================================
    print("[5/5] Renderizando Diagramas Mermaid y Capítulos X al XIV...")
    p_c10 = doc.add_paragraph()
    p_c10.paragraph_format.space_before = Pt(16)
    p_c10.paragraph_format.space_after = Pt(12)
    p_c10.paragraph_format.keep_with_next = True
    r_c10 = p_c10.add_run(part3.CAPITULO_X["titulo"])
    r_c10.font.name = "Arial"
    r_c10.font.size = Pt(14)
    r_c10.bold = True
    r_c10.font.color.rgb = RGBColor(30, 58, 138)

    # Explicación y Justificación de la Metodología en Cascada
    sec_cascada = part3.CAPITULO_X["secciones"][0]
    p_sc = doc.add_paragraph()
    p_sc.paragraph_format.space_before = Pt(14)
    p_sc.paragraph_format.space_after = Pt(6)
    p_sc.paragraph_format.keep_with_next = True
    r_sc = p_sc.add_run(sec_cascada["subtitulo"])
    r_sc.font.name = "Arial"
    r_sc.font.size = Pt(12.5)
    r_sc.bold = True
    r_sc.font.color.rgb = RGBColor(15, 23, 42)

    for parrafo in sec_cascada["contenido"].split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p, parrafo.strip())

    # Diagrama 10.0: Ciclo de Vida en Cascada
    render_mermaid_diagram(doc, part3.MERMAID_DIAGRAMS[0])

    # 10.1 Análisis
    sec_10_1 = part3.CAPITULO_X["secciones"][1]
    p_s101 = doc.add_paragraph()
    p_s101.paragraph_format.space_before = Pt(14)
    p_s101.paragraph_format.space_after = Pt(6)
    p_s101.paragraph_format.keep_with_next = True
    r_s101 = p_s101.add_run(sec_10_1["subtitulo"])
    r_s101.font.name = "Arial"
    r_s101.font.size = Pt(12.5)
    r_s101.bold = True
    r_s101.font.color.rgb = RGBColor(15, 23, 42)

    for parrafo in sec_10_1["contenido"].split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p, parrafo.strip())

    # Diagrama 10.1: Casos de Uso
    render_mermaid_diagram(doc, part3.MERMAID_DIAGRAMS[1])

    # 10.2 Diseño con Diagramas Mermaid
    sec_10_2 = part3.CAPITULO_X["secciones"][2]
    p_s102 = doc.add_paragraph()
    p_s102.paragraph_format.space_before = Pt(14)
    p_s102.paragraph_format.space_after = Pt(6)
    p_s102.paragraph_format.keep_with_next = True
    r_s102 = p_s102.add_run(sec_10_2["subtitulo"])
    r_s102.font.name = "Arial"
    r_s102.font.size = Pt(12.5)
    r_s102.bold = True
    r_s102.font.color.rgb = RGBColor(15, 23, 42)

    p_d102 = doc.add_paragraph()
    p_d102.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_formatted_text(p_d102, sec_10_2["contenido"])

    # Renderizar restantes diagramas Mermaid (C4, DER, Secuencia, Estados)
    for diag in part3.MERMAID_DIAGRAMS[2:]:
        render_mermaid_diagram(doc, diag)

    # 10.3 a 10.7
    for sub in part3.CAPITULO_X_CONTINUACION["subsecciones"]:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(12)
        p_sub.paragraph_format.space_after = Pt(6)
        p_sub.paragraph_format.keep_with_next = True
        r_sub = p_sub.add_run(sub["subtitulo"])
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(12)
        r_sub.bold = True
        r_sub.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sub["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

        # Si es la sección de recursos 10.7
        if "10.7 RECURSOS" in sub["subtitulo"]:
            rec = part3.RECURSOS_DETALLE
            
            p_mat = doc.add_paragraph()
            p_mat.paragraph_format.space_before = Pt(8)
            p_mat.paragraph_format.space_after = Pt(4)
            r_mat = p_mat.add_run("10.7.1 RECURSOS MATERIALES")
            r_mat.bold = True
            for parrafo in rec["materiales"].split("\n"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline_formatted_text(p, parrafo.strip())

            p_hum = doc.add_paragraph()
            p_hum.paragraph_format.space_before = Pt(8)
            p_hum.paragraph_format.space_after = Pt(4)
            r_hum = p_hum.add_run("10.7.2 RECURSOS HUMANOS")
            r_hum.bold = True
            for parrafo in rec["humanos"].split("\n"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline_formatted_text(p, parrafo.strip())

            p_pre = doc.add_paragraph()
            p_pre.paragraph_format.space_before = Pt(8)
            p_pre.paragraph_format.space_after = Pt(4)
            r_pre = p_pre.add_run("10.7.3 PRESUPUESTO")
            r_pre.bold = True
            render_table_in_doc(doc, rec["presupuesto_tabla"])

    doc.add_page_break()

    # ==========================================================================
    # CAPÍTULOS XI AL XIV
    # ==========================================================================
    # Cap XI: Articulación
    p_c11 = doc.add_paragraph()
    p_c11.paragraph_format.space_before = Pt(16)
    p_c11.paragraph_format.space_after = Pt(12)
    p_c11.paragraph_format.keep_with_next = True
    r_c11 = p_c11.add_run(part3.CAPITULO_XI["titulo"])
    r_c11.font.name = "Arial"
    r_c11.font.size = Pt(14)
    r_c11.bold = True
    r_c11.font.color.rgb = RGBColor(30, 58, 138)

    p_d11 = doc.add_paragraph()
    p_d11.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_formatted_text(p_d11, part3.CAPITULO_XI["descripcion"])

    for campo in part3.CAPITULO_XI["campos"]:
        p_cp = doc.add_paragraph()
        p_cp.paragraph_format.space_before = Pt(12)
        p_cp.paragraph_format.space_after = Pt(4)
        p_cp.paragraph_format.keep_with_next = True
        r_cp = p_cp.add_run(campo["nombre"])
        r_cp.font.name = "Arial"
        r_cp.font.size = Pt(12.5)
        r_cp.bold = True
        r_cp.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in campo["articulacion"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap XII: Conclusiones y Recomendaciones
    p_c12 = doc.add_paragraph()
    p_c12.paragraph_format.space_before = Pt(16)
    p_c12.paragraph_format.space_after = Pt(12)
    p_c12.paragraph_format.keep_with_next = True
    r_c12 = p_c12.add_run(part3.CAPITULO_XII["titulo"])
    r_c12.font.name = "Arial"
    r_c12.font.size = Pt(14)
    r_c12.bold = True
    r_c12.font.color.rgb = RGBColor(30, 58, 138)

    for sec in part3.CAPITULO_XII["secciones"]:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.keep_with_next = True
        r_s = p_s.add_run(sec["subtitulo"])
        r_s.font.name = "Arial"
        r_s.font.size = Pt(12.5)
        r_s.bold = True
        r_s.font.color.rgb = RGBColor(15, 23, 42)

        for parrafo in sec["contenido"].split("\n\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap XIII: Proyecto de Vida
    p_c13 = doc.add_paragraph()
    p_c13.paragraph_format.space_before = Pt(16)
    p_c13.paragraph_format.space_after = Pt(12)
    p_c13.paragraph_format.keep_with_next = True
    r_c13 = p_c13.add_run(part3.CAPITULO_XIII["titulo"])
    r_c13.font.name = "Arial"
    r_c13.font.size = Pt(14)
    r_c13.bold = True
    r_c13.font.color.rgb = RGBColor(30, 58, 138)

    for parrafo in part3.CAPITULO_XIII["contenido"].split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p, parrafo.strip())

    doc.add_page_break()

    # Cap XIV: Bibliografía
    p_c14 = doc.add_paragraph()
    p_c14.paragraph_format.space_before = Pt(16)
    p_c14.paragraph_format.space_after = Pt(12)
    p_c14.paragraph_format.keep_with_next = True
    r_c14 = p_c14.add_run(part3.CAPITULO_XIV["titulo"])
    r_c14.font.name = "Arial"
    r_c14.font.size = Pt(14)
    r_c14.bold = True
    r_c14.font.color.rgb = RGBColor(30, 58, 138)

    for ref in part3.CAPITULO_XIV["referencias"]:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.15
        add_inline_formatted_text(p_ref, ref)

    doc.add_page_break()

    # ==========================================================================
    # ANEXOS
    # ==========================================================================
    p_anx = doc.add_paragraph()
    p_anx.paragraph_format.space_before = Pt(16)
    p_anx.paragraph_format.space_after = Pt(12)
    p_anx.paragraph_format.keep_with_next = True
    r_anx = p_anx.add_run(part3.ANEXOS["titulo"])
    r_anx.font.name = "Arial"
    r_anx.font.size = Pt(14)
    r_anx.bold = True
    r_anx.font.color.rgb = RGBColor(30, 58, 138)

    for anexo in part3.ANEXOS["secciones"]:
        p_sa = doc.add_paragraph()
        p_sa.paragraph_format.space_before = Pt(14)
        p_sa.paragraph_format.space_after = Pt(6)
        p_sa.paragraph_format.keep_with_next = True
        r_sa = p_sa.add_run(anexo["subtitulo"])
        r_sa.font.name = "Arial"
        r_sa.font.size = Pt(12.5)
        r_sa.bold = True
        r_sa.font.color.rgb = RGBColor(15, 23, 42)

        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatted_text(p_desc, anexo["descripcion"])

        if "tablas" in anexo:
            for t_item in anexo["tablas"]:
                render_table_in_doc(doc, {
                    "titulo": t_item["nombre"],
                    "columnas": t_item["columnas"],
                    "filas": t_item["filas"]
                })
        
        if "tabla" in anexo:
            render_table_in_doc(doc, anexo["tabla"])

        if "codigo_haversine" in anexo:
            render_code_block(doc, anexo["codigo_haversine"], label="ALGORITMO GEODÉSICO HAVERSINE (Python 3.11)")

        if "codigo_for_update" in anexo:
            render_code_block(doc, anexo["codigo_for_update"], label="BLOQUEO PESIMISTA TRANSACCIONAL ACID (PHP 8.2 PDO)")

    # Guardar documento Word
    doc.save(OUT_DOCX)
    print(f"[OK] Documento Word generado exitosamente: {OUT_DOCX}")

    # Copiar/actualizar en Downloads si existe la ruta
    try:
        if os.path.exists(os.path.dirname(DOWNLOADS_DOCX)):
            if os.path.exists(DOWNLOADS_DOCX) and not os.path.exists(DOWNLOADS_DOCX + ".bak"):
                shutil.copy2(DOWNLOADS_DOCX, DOWNLOADS_DOCX + ".bak")
                print(f"[BACKUP] Respaldo creado en: {DOWNLOADS_DOCX}.bak")
            shutil.copy2(OUT_DOCX, DOWNLOADS_DOCX)
            print(f"[OK] Archivo actualizado directamente en Downloads: {DOWNLOADS_DOCX}")
    except Exception as e:
        print(f"[WARN] No se pudo copiar a Downloads: {e}")

# ==============================================================================
# GENERADOR HTML COMPANION CON MERMAID.JS INTERACTIVO
# ==============================================================================

def build_monografia_html():
    """Genera una versión HTML interactiva con Mermaid.js para renderizado en navegadores."""
    print("Compilando visor interactivo HTML con diagramas vectoriales Mermaid.js...")
    meta = part1.METADATA
    
    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Monografía BTH 2026 - Burger 24/7 E-Commerce</title>
    <script type="module">
        import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.esm.min.mjs';
        mermaid.initialize({{
            startOnLoad: true,
            theme: 'default',
            securityLevel: 'loose',
            fontFamily: 'Arial, sans-serif'
        }});
    </script>
    <style>
        :root {{
            --primary: #1e3a8a;
            --primary-dark: #0f172a;
            --secondary: #d97706;
            --text: #1f2937;
            --bg: #f8fafc;
            --card-bg: #ffffff;
            --border: #e2e8f0;
        }}
        body {{
            font-family: 'Arial', sans-serif;
            line-height: 1.6;
            color: var(--text);
            background-color: var(--bg);
            margin: 0;
            padding: 0;
        }}
        .container {{
            max-width: 900px;
            margin: 0 auto;
            background: var(--card-bg);
            padding: 50px 70px;
            box-shadow: 0 4px 20px rgba(0,0,0,0.06);
        }}
        .cover {{
            text-align: center;
            border-bottom: 2px solid var(--border);
            padding-bottom: 40px;
            margin-bottom: 40px;
        }}
        .cover h1 {{
            color: var(--primary-dark);
            font-size: 1.6rem;
            margin-bottom: 5px;
        }}
        .cover h2 {{
            color: var(--primary);
            font-size: 1.8rem;
            margin-top: 0;
            margin-bottom: 10px;
        }}
        .cover .subtitle {{
            color: var(--secondary);
            font-weight: bold;
            font-size: 1rem;
            text-transform: uppercase;
            letter-spacing: 1px;
        }}
        .cover .title-box {{
            margin: 40px 0;
            padding: 20px;
            background: #eff6ff;
            border-left: 6px solid var(--primary);
            font-size: 1.4rem;
            font-weight: bold;
            color: var(--primary-dark);
            text-align: center;
        }}
        .cover .meta-box {{
            font-size: 1.05rem;
            line-height: 1.8;
            margin-top: 30px;
        }}
        h2.chap-title {{
            color: var(--primary);
            border-bottom: 2px solid #bfdbfe;
            padding-bottom: 8px;
            margin-top: 40px;
            font-size: 1.35rem;
        }}
        h3.sec-title {{
            color: var(--primary-dark);
            margin-top: 25px;
            font-size: 1.15rem;
        }}
        p {{
            text-align: justify;
            margin-bottom: 14px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            font-size: 0.95rem;
        }}
        table th {{
            background-color: var(--primary);
            color: #ffffff;
            padding: 10px;
            text-align: left;
        }}
        table td {{
            padding: 9px 10px;
            border-bottom: 1px solid var(--border);
        }}
        table tr:nth-child(even) {{
            background-color: #f8fafc;
        }}
        .mermaid-container {{
            background: #ffffff;
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 20px;
            margin: 25px 0;
            box-shadow: 0 2px 8px rgba(0,0,0,0.04);
            text-align: center;
        }}
        .mermaid-title {{
            font-weight: bold;
            color: var(--primary-dark);
            margin-bottom: 8px;
            font-size: 1.05rem;
        }}
        .mermaid-desc {{
            font-size: 0.9rem;
            color: #64748b;
            font-style: italic;
            margin-bottom: 15px;
        }}
        pre.code-block {{
            background: #1e293b;
            color: #f8fafc;
            padding: 15px;
            border-radius: 6px;
            overflow-x: auto;
            font-family: 'Consolas', monospace;
            font-size: 0.88rem;
        }}
        .print-btn {{
            position: fixed;
            top: 20px;
            right: 20px;
            background: var(--primary);
            color: white;
            padding: 10px 18px;
            border: none;
            border-radius: 6px;
            cursor: pointer;
            font-weight: bold;
            box-shadow: 0 4px 10px rgba(0,0,0,0.15);
        }}
        @media print {{
            .print-btn {{ display: none; }}
            .container {{ box-shadow: none; padding: 0; }}
            body {{ background: white; }}
        }}
    </style>
</head>
<body>
    <button class="print-btn" onclick="window.print()">🖨️ Imprimir / Guardar PDF</button>
    <div class="container">
        <div class="cover">
            <h1>{meta['institucion']}</h1>
            <h2>{meta['colegio'].replace(chr(10), '<br>')}</h2>
            <div class="subtitle">{meta['especialidad']} - {meta['subtitulo_bth']}</div>
            <div class="title-box">{meta['titulo']}</div>
            <div class="meta-box">
                <strong>Postulantes:</strong> {meta.get('postulantes', meta.get('postulante', 'Nataly Gemio'))}<br>
                <strong>Tutor:</strong> {meta['tutor']}<br>
                <strong>Curso:</strong> {meta['curso']}<br>
                <strong>{meta['lugar']} &bull; {meta['gestion']}</strong>
            </div>
        </div>

        <h2 class="chap-title">I. INTRODUCCIÓN</h2>
        {"".join(f"<p>{html.escape(p)}</p>" for p in part1.CAPITULO_I['contenido'].split(chr(10)+chr(10)) if p.strip())}

        <h2 class="chap-title">VIII. MARCO PROCEDIMENTAL</h2>
        <h3 class="sec-title">8.1 PROPUESTA DE INNOVACIÓN</h3>
        {"".join(f"<p>{html.escape(p)}</p>" for p in part3.CAPITULO_VIII['secciones'][0]['contenido'].split(chr(10)+chr(10)) if p.strip())}
        <h3 class="sec-title">8.2 RESULTADOS ESPERADOS</h3>
        {"".join(f"<p>{html.escape(p)}</p>" for p in part3.CAPITULO_VIII['secciones'][1]['contenido'].split(chr(10)+chr(10)) if p.strip())}

        <h2 class="chap-title">VII. MARCO TEÓRICO</h2>
        <h3 class="sec-title">7.1 SUSTENTO LEGAL</h3>
        <p>{html.escape(part2.CAPITULO_VII['seccion_7_1']['descripcion'])}</p>
"""
    # Agregar items legales
    for item in part2.CAPITULO_VII['seccion_7_1']['items']:
        html_content += f"<h4 style='color:#1e3a8a;margin-top:15px;'>{item['codigo']} {html.escape(item['nombre'])}</h4>"
        for p in item['analisis'].split("\n\n"):
            if p.strip():
                html_content += f"<p>{html.escape(p.strip())}</p>"

    # Agregar temas 7.2
    html_content += f"<h3 class='sec-title'>7.2 DESARROLLO DEL MARCO TEÓRICO</h3>"
    html_content += f"<p>{html.escape(part2.CAPITULO_VII['seccion_7_2']['descripcion'])}</p>"
    for tema in part2.CAPITULO_VII['seccion_7_2']['temas']:
        html_content += f"<h4 style='color:#0f172a;margin-top:20px;font-size:1.1rem;'>{tema['numero']} {html.escape(tema['titulo'])}</h4>"
        for p in tema['contenido'].split("\n\n"):
            if p.strip():
                html_content += f"<p>{html.escape(p.strip())}</p>"

    # Agregar Diagramas Mermaid interactivos
    html_content += "<h2 class='chap-title'>X. DIAGRAMAS Y MODELADO DEL SISTEMA (MERMAID.JS)</h2>"
    for diag in part3.MERMAID_DIAGRAMS:
        html_content += f"""
        <div class="mermaid-container">
            <div class="mermaid-title">Figura {diag['numero']}: {html.escape(diag['titulo'])}</div>
            <div class="mermaid-desc">{html.escape(diag['descripcion'])}</div>
            <div class="mermaid">
{diag['mermaid']}
            </div>
        </div>
        """

    html_content += """
    </div>
</body>
</html>
"""
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"[OK] Visor interactivo HTML generado: {OUT_HTML}")

# ==============================================================================
# PUNTO DE ENTRADA
# ==============================================================================
if __name__ == "__main__":
    build_monografia_docx()
    build_monografia_html()
    print("\n======================================================================")
    print("  COMPILACIÓN DE MONOGRAFÍA BTH 2026 COMPLETADA SATISFACTORIAMENTE")
    print("======================================================================")
